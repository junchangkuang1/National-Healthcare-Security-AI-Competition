"""Dataset discovery, validation, clinical-label parsing, and tensor loading."""

from __future__ import annotations

import gzip
import json
import random
import re
import sys
import unicodedata
from dataclasses import asdict, dataclass
from functools import lru_cache
from pathlib import Path
from typing import Any, Optional, Sequence

import cv2
import nibabel as nib
import numpy as np
import openpyxl
import pydicom
import torch
from pydicom.pixel_data_handlers.util import apply_voi_lut
from torch.utils.data import Dataset

IGNORE_INDEX = -100
IMAGENET_MEAN = np.asarray([0.485, 0.456, 0.406], dtype=np.float32)
IMAGENET_STD = np.asarray([0.229, 0.224, 0.225], dtype=np.float32)
MASK_HINTS = ("_mask", "mask_", "病灶", "主病灶")


@dataclass
class Sample:
    patient_id: str
    kind: str
    image_path: str
    mask_path: str
    frame_index: int
    series_uid: str
    accession_number: str
    roi_name: str
    detection: int
    malignancy: int = IGNORE_INDEX
    tirads: int = IGNORE_INDEX
    extension: int = IGNORE_INDEX
    location: int = IGNORE_INDEX
    t_clinsig: int = IGNORE_INDEX


class DataReadError(RuntimeError):
    """Error that preserves the path and format of an unreadable medical file."""


def data_read_error(kind: str, path: Any, exc: BaseException) -> DataReadError:
    return DataReadError(
        f"Failed to read {kind} file: {path}\n"
        f"Original error: {type(exc).__name__}: {exc}\n"
        "The file may be truncated, corrupt, or use an unsupported compression/transfer syntax."
    )


def norm_text(value: Any) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value)).strip().lower()
    return re.sub(r"[\s_\-—:：>/\\().\[\]]+", "", text)


def first_int(value: Any) -> Optional[int]:
    if value is None or str(value).strip() == "":
        return None
    match = re.search(r"-?\d+", unicodedata.normalize("NFKC", str(value)))
    return int(match.group()) if match else None


def parse_binary(value: Any, positive: Sequence[str], negative: Sequence[str]) -> int:
    text = norm_text(value)
    if not text or text in {"null", "nan", "na", "n/a", "unk", "无数据"}:
        return IGNORE_INDEX
    if any(norm_text(item) == text or norm_text(item) in text for item in positive):
        return 1
    if any(norm_text(item) == text or norm_text(item) in text for item in negative):
        return 0
    number = first_int(value)
    return number if number in (0, 1) else IGNORE_INDEX


def parse_class(value: Any, minimum: int, maximum: int) -> int:
    if value is None:
        return IGNORE_INDEX
    text = unicodedata.normalize("NFKC", str(value)).upper().strip()
    if not text or text in {"NULL", "NAN", "NA", "N/A", "UNK", "无"}:
        return IGNORE_INDEX
    roman = {"I": 1, "II": 2, "III": 3, "IV": 4, "V": 5}
    number = roman.get(re.sub(r"[^IVX]", "", text))
    if number is None:
        number = first_int(text)
    if number is None or not minimum <= number <= maximum:
        return IGNORE_INDEX
    return number - minimum


def normalize_roi_name(name: Any) -> str:
    text = norm_text(name)
    if not text:
        return ""
    if text.startswith("m") or "主病灶" in text:
        return "m"
    match = re.search(r"病灶(\d+)", text)
    return f"病灶{match.group(1)}" if match else text


def roi_name_from_mask(path: Path) -> str:
    name = path.name[:-7] if path.name.lower().endswith(".nii.gz") else path.stem
    name = re.sub(r"(?i)_?mask.*$", "", name).strip("_- ")
    if re.match(r"(?i)^m(?:_|\d|$)", name):
        return "M"
    match = re.search(r"病灶\s*(\d+)", name)
    return f"病灶{match.group(1)}" if match else name


class ClinicalTable:
    """Read the competition spreadsheet, including its documented two-row layout."""

    ALIASES = {
        "patient": ("patientid", "patientld", "researchid"),
        "series": ("seriesuid", "seriesld"),
        "accession": ("accessionnumber",),
        "roi": ("roiname",),
        "malignancy": ("roiClinical良恶性", "良恶性"),
        "tirads": ("roiClinicalTIRADS", "acrtirads分类", "tirads"),
        "extension": ("甲状腺外侵犯", "extrathyext"),
        "location": ("主病灶位置", "lesionloc", "esionloc"),
        "t_clinsig": ("临床t分期", "临床显著性分级t", "clintstage"),
    }

    def __init__(self, xlsx_paths: Sequence[Path]):
        self.rows: list[dict[str, Any]] = []
        for path in xlsx_paths:
            self._read_workbook(path)
        self.indices: dict[str, dict[tuple[str, ...], list[dict[str, Any]]]] = {
            name: {} for name in ("psr", "par", "pr", "sr", "ar", "r")
        }
        for row in self.rows:
            patient = norm_text(row.get("patient"))
            series = norm_text(row.get("series"))
            accession = norm_text(row.get("accession"))
            roi = normalize_roi_name(row.get("roi"))
            keys = (
                ("psr", (patient, series, roi)), ("par", (patient, accession, roi)),
                ("pr", (patient, roi)), ("sr", (series, roi)),
                ("ar", (accession, roi)), ("r", (roi,)),
            )
            for index_name, key in keys:
                if all(key):
                    self.indices[index_name].setdefault(key, []).append(row)

    @classmethod
    def _field(cls, headers: Sequence[Any], field: str) -> Optional[int]:
        normalized = [norm_text(header) for header in headers]
        aliases = [norm_text(alias) for alias in cls.ALIASES[field]]
        for alias in aliases:
            for index, header in enumerate(normalized):
                if header == alias:
                    return index
        for alias in aliases:
            for index, header in enumerate(normalized):
                if alias and alias in header:
                    return index
        return None

    def _read_workbook(self, path: Path) -> None:
        workbook = None
        try:
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
            for worksheet in workbook.worksheets:
                preview = list(
                    worksheet.iter_rows(min_row=1, max_row=min(20, worksheet.max_row), values_only=True)
                )
                header_row, headers = None, ()
                for row_number, row in enumerate(preview, start=1):
                    normalized = {norm_text(value) for value in row if value is not None}
                    has_patient = any(value in normalized for value in ("patientid", "patientld", "researchid"))
                    has_exam = any(value in normalized for value in ("seriesuid", "seriesld", "accessionnumber"))
                    if has_patient and has_exam:
                        header_row, headers = row_number, row
                        break
                if header_row is None:
                    continue
                columns = {field: self._field(headers, field) for field in self.ALIASES}
                for values in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
                    if not any(value is not None and str(value).strip() for value in values):
                        continue
                    item = {
                        field: values[index] if index is not None and index < len(values) else None
                        for field, index in columns.items()
                    }
                    if item.get("patient") or item.get("series") or item.get("accession"):
                        self.rows.append(item)
        except Exception as exc:
            raise data_read_error("Excel workbook", path, exc) from exc
        finally:
            if workbook is not None:
                workbook.close()

    def lookup(self, patient: str, series: str, accession: str, roi: str) -> dict[str, int]:
        patient_key, series_key = norm_text(patient), norm_text(series)
        accession_key, roi_key = norm_text(accession), normalize_roi_name(roi)
        candidates: list[dict[str, Any]] = []
        for index_name, key in (
            ("psr", (patient_key, series_key, roi_key)), ("par", (patient_key, accession_key, roi_key)),
            ("pr", (patient_key, roi_key)), ("sr", (series_key, roi_key)),
            ("ar", (accession_key, roi_key)), ("r", (roi_key,)),
        ):
            if all(key) and key in self.indices[index_name]:
                candidates = self.indices[index_name][key]
                break
        row = candidates[0] if candidates else {}
        return {
            "malignancy": parse_binary(row.get("malignancy"), ("恶", "恶性", "阳性"), ("良", "良性", "阴性")),
            "tirads": parse_class(row.get("tirads"), 1, 5),
            "extension": parse_class(row.get("extension"), 1, 4),
            "location": parse_class(row.get("location"), 1, 7),
            "t_clinsig": parse_binary(row.get("t_clinsig"), ("有", "阳性"), ("无", "阴性")),
        }


def is_mask_path(path: Path) -> bool:
    name = path.name.lower()
    return any(hint in name for hint in MASK_HINTS) or bool(re.match(r"^m(?:_|\d)", name))


def patient_from_path(path: Path, root: Path) -> str:
    try:
        relative = path.relative_to(root)
        return relative.parts[0] if relative.parts else path.parent.name
    except ValueError:
        return path.parent.name


def dicom_metadata(path: Path) -> tuple[str, str, str]:
    try:
        dataset = pydicom.dcmread(path, stop_before_pixels=True, force=True)
        patient = str(getattr(dataset, "PatientID", "") or "")
        series = str(getattr(dataset, "SeriesInstanceUID", "") or path.parent.name)
        accession = str(getattr(dataset, "AccessionNumber", "") or "")
        return patient, series, accession
    except Exception:
        return "", path.parent.name, ""


@lru_cache(maxsize=4)
def load_nifti_array(path: str) -> np.ndarray:
    try:
        image = nib.load(path)
        return np.asanyarray(image.dataobj)
    except Exception as exc:
        raise data_read_error("NIfTI", path, exc) from exc


def positive_frames(mask_path: Path) -> list[int]:
    array = np.squeeze(load_nifti_array(str(mask_path)))
    if array.ndim <= 2:
        return [0] if np.any(array > 0) else []
    if array.ndim != 3:
        raise ValueError(f"Unsupported mask dimensions {array.shape}: {mask_path}")
    return np.flatnonzero(np.any(array > 0, axis=(0, 1))).astype(int).tolist()


def nifti_frame_count(path: Path) -> int:
    try:
        shape = nib.load(str(path)).shape
        return int(shape[2]) if len(shape) >= 3 else 1
    except Exception as exc:
        raise data_read_error("NIfTI header", path, exc) from exc


@lru_cache(maxsize=None)
def nifti_validation_error(path: str) -> Optional[str]:
    try:
        if path.lower().endswith(".gz"):
            with gzip.open(path, "rb") as stream:
                while stream.read(8 * 1024 * 1024):
                    pass
        image = nib.load(path)
        if not image.shape or any(int(dimension) <= 0 for dimension in image.shape):
            raise ValueError(f"invalid NIfTI shape: {image.shape}")
        return None
    except Exception as exc:
        return f"{type(exc).__name__}: {exc}"


def build_samples(
    root: Path,
    clinical: ClinicalTable,
    negative_ratio: float,
    seed: int,
    bad_file_policy: str,
) -> tuple[list[Sample], dict[str, str]]:
    if not root.is_dir():
        raise FileNotFoundError(f"Training annotation directory not found: {root}")
    samples: list[Sample] = []
    bad_files: dict[str, str] = {}
    rng = random.Random(seed)

    def usable_nifti(path: Path) -> bool:
        error = nifti_validation_error(str(path))
        if error is None:
            return True
        bad_files[str(path)] = error
        message = f"[bad-nifti] {path} | {error}"
        if bad_file_policy == "error":
            raise DataReadError(message)
        print(message + " | skipped", file=sys.stderr, flush=True)
        return False

    for dicom_path in sorted(root.rglob("*.dcm")):
        candidate_masks = sorted(
            path for path in dicom_path.parent.glob("*.nii*") if path.is_file() and is_mask_path(path)
        )
        masks = [path for path in candidate_masks if usable_nifti(path)]
        dicom_patient, series_uid, accession = dicom_metadata(dicom_path)
        patient = dicom_patient or patient_from_path(dicom_path, root)
        if candidate_masks and not masks:
            print(f"[skip-series] all masks are unreadable: {dicom_path.parent}", file=sys.stderr, flush=True)
            continue
        if not masks:
            samples.append(Sample(patient, "static", str(dicom_path), "", -1, series_uid, accession, "", 0))
            continue
        for mask in masks:
            roi = roi_name_from_mask(mask)
            labels = clinical.lookup(patient, series_uid, accession, roi)
            samples.append(
                Sample(patient, "static", str(dicom_path), str(mask), -1, series_uid, accession, roi, 1, **labels)
            )

    all_nifti = sorted({*root.rglob("*.nii"), *root.rglob("*.nii.gz")})
    raw_volumes = [path for path in all_nifti if path.is_file() and not is_mask_path(path)]
    for volume in raw_volumes:
        if not usable_nifti(volume):
            print(f"[skip-volume] unreadable image: {volume}", file=sys.stderr, flush=True)
            continue
        candidate_masks = sorted(
            path for path in volume.parent.glob("*.nii*") if path != volume and is_mask_path(path)
        )
        masks = [path for path in candidate_masks if usable_nifti(path)]
        if candidate_masks and not masks:
            print(f"[skip-volume] all masks are unreadable: {volume.parent}", file=sys.stderr, flush=True)
            continue
        if not masks:
            continue
        patient = patient_from_path(volume, root)
        accession = volume.parent.name
        frame_total = nifti_frame_count(volume)
        all_positive: set[int] = set()
        for mask in masks:
            roi = roi_name_from_mask(mask)
            labels = clinical.lookup(patient, "", accession, roi)
            frames = [frame for frame in positive_frames(mask) if 0 <= frame < frame_total]
            all_positive.update(frames)
            for frame in frames:
                samples.append(
                    Sample(patient, "dynamic", str(volume), str(mask), frame, "", accession, roi, 1, **labels)
                )
        negative_frames = [frame for frame in range(frame_total) if frame not in all_positive]
        requested = min(len(negative_frames), max(1, round(len(all_positive) * negative_ratio)))
        for frame in rng.sample(negative_frames, requested) if requested else []:
            samples.append(Sample(patient, "dynamic", str(volume), "", frame, "", accession, "", 0))

    if not samples:
        raise RuntimeError("No samples were indexed. Check the directory structure and NIfTI mask naming.")
    return samples, bad_files


def split_by_patient(
    samples: Sequence[Sample], seed: int
) -> tuple[list[Sample], list[Sample], list[Sample], dict[str, list[str]]]:
    patients = sorted({sample.patient_id for sample in samples})
    if len(patients) < 3:
        raise RuntimeError("At least three distinct patient identifiers are required for an 8:1:1 split")
    rng = random.Random(seed)
    rng.shuffle(patients)
    validation_count = max(1, round(len(patients) * 0.1))
    test_count = max(1, round(len(patients) * 0.1))
    if validation_count + test_count >= len(patients):
        validation_count = test_count = 1
    validation_ids = set(patients[:validation_count])
    test_ids = set(patients[validation_count:validation_count + test_count])
    train_ids = set(patients[validation_count + test_count:])
    groups = {
        "train": sorted(train_ids), "val": sorted(validation_ids), "test": sorted(test_ids),
    }
    return (
        [sample for sample in samples if sample.patient_id in train_ids],
        [sample for sample in samples if sample.patient_id in validation_ids],
        [sample for sample in samples if sample.patient_id in test_ids],
        groups,
    )


def save_samples(path: Path, samples: Sequence[Sample]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="\n") as stream:
        for sample in samples:
            stream.write(json.dumps(asdict(sample), ensure_ascii=False, separators=(",", ":")) + "\n")


def load_samples(path: Path) -> list[Sample]:
    with path.open("r", encoding="utf-8") as stream:
        return [Sample(**json.loads(line)) for line in stream if line.strip()]


def select_split(samples: Sequence[Sample], split_ids: dict[str, list[str]], name: str) -> list[Sample]:
    patient_ids = set(split_ids[name])
    return [sample for sample in samples if sample.patient_id in patient_ids]


def read_dicom_image(path: str) -> np.ndarray:
    try:
        dataset = pydicom.dcmread(path, force=True)
        array = dataset.pixel_array
    except Exception as exc:
        raise data_read_error("DICOM pixel data", path, exc) from exc
    try:
        array = apply_voi_lut(array, dataset)
    except Exception:
        pass
    array = np.squeeze(np.asarray(array))
    if array.ndim == 3:
        array = array[0]
    if array.ndim != 2:
        raise ValueError(f"Unsupported DICOM pixel shape {array.shape}: {path}")
    if str(getattr(dataset, "PhotometricInterpretation", "")) == "MONOCHROME1":
        array = array.max() - array
    return array.astype(np.float32)


def read_nifti_slice(path: str, frame: int) -> np.ndarray:
    array = np.squeeze(load_nifti_array(path))
    if array.ndim == 2:
        plane = array
    elif array.ndim == 3:
        if not 0 <= frame < array.shape[2]:
            raise IndexError(f"Frame {frame} outside shape {array.shape}: {path}")
        plane = array[:, :, frame]
    else:
        raise ValueError(f"Unsupported NIfTI shape {array.shape}: {path}")
    return np.asarray(plane).T


def robust_normalize(image: np.ndarray) -> np.ndarray:
    image = np.nan_to_num(image.astype(np.float32), nan=0.0, posinf=0.0, neginf=0.0)
    finite = image[np.isfinite(image)]
    if finite.size == 0:
        return np.zeros_like(image, dtype=np.float32)
    low, high = np.percentile(finite, (0.5, 99.5))
    if high <= low:
        low, high = float(image.min()), float(image.max())
    if high <= low:
        return np.zeros_like(image, dtype=np.float32)
    return np.clip((image - low) / (high - low), 0.0, 1.0)


def align_mask(mask: np.ndarray, image_shape: tuple[int, int]) -> np.ndarray:
    mask = np.squeeze(mask)
    if mask.ndim == 3:
        mask = mask[:, :, 0]
    if mask.ndim != 2:
        raise ValueError(f"Mask must become 2D, got {mask.shape}")
    for candidate in (mask, mask.T):
        if candidate.shape == image_shape:
            return (candidate > 0).astype(np.float32)
    resized = cv2.resize(mask.astype(np.float32), (image_shape[1], image_shape[0]), interpolation=cv2.INTER_NEAREST)
    return (resized > 0).astype(np.float32)


def image_to_tensor(image: np.ndarray, image_size: int) -> torch.Tensor:
    normalized = robust_normalize(image)
    resized = cv2.resize(normalized, (image_size, image_size), interpolation=cv2.INTER_LINEAR)
    rgb = np.repeat(resized[..., None], 3, axis=2)
    rgb = (rgb - IMAGENET_MEAN) / IMAGENET_STD
    return torch.from_numpy(rgb.transpose(2, 0, 1).astype(np.float32))


class ThyroidDataset(Dataset):
    def __init__(self, samples: Sequence[Sample], image_size: int, training: bool):
        self.samples = list(samples)
        self.image_size = image_size
        self.training = training

    def __len__(self) -> int:
        return len(self.samples)

    def __getitem__(self, index: int) -> dict[str, torch.Tensor]:
        sample = self.samples[index]
        if sample.kind == "static":
            image = read_dicom_image(sample.image_path)
            raw_mask = np.squeeze(load_nifti_array(sample.mask_path)) if sample.mask_path else np.zeros(image.shape)
            mask = align_mask(raw_mask, image.shape)
        else:
            image = read_nifti_slice(sample.image_path, sample.frame_index).astype(np.float32)
            raw_mask = (
                read_nifti_slice(sample.mask_path, sample.frame_index)
                if sample.mask_path else np.zeros(image.shape, dtype=np.float32)
            )
            mask = align_mask(raw_mask, image.shape)

        image = robust_normalize(image)
        if self.training:
            if random.random() < 0.5:
                image, mask = np.fliplr(image).copy(), np.fliplr(mask).copy()
            if random.random() < 0.15:
                image, mask = np.flipud(image).copy(), np.flipud(mask).copy()
            if random.random() < 0.35:
                rotations = random.randint(1, 3)
                image, mask = np.rot90(image, rotations).copy(), np.rot90(mask, rotations).copy()
            if random.random() < 0.35:
                image = np.clip(image * random.uniform(0.85, 1.15) + random.uniform(-0.08, 0.08), 0.0, 1.0)
            if random.random() < 0.20:
                noise = np.random.normal(0, 0.02, image.shape).astype(np.float32)
                image = np.clip(image + noise, 0.0, 1.0)

        size = (self.image_size, self.image_size)
        image = cv2.resize(image, size, interpolation=cv2.INTER_LINEAR)
        mask = cv2.resize(mask, size, interpolation=cv2.INTER_NEAREST)
        rgb = np.repeat(image[..., None], 3, axis=2)
        rgb = (rgb - IMAGENET_MEAN) / IMAGENET_STD
        return {
            "image": torch.from_numpy(rgb.transpose(2, 0, 1).astype(np.float32)),
            "mask": torch.from_numpy((mask > 0.5).astype(np.float32)[None]),
            "detection": torch.tensor(sample.detection, dtype=torch.float32),
            "malignancy": torch.tensor(sample.malignancy, dtype=torch.long),
            "tirads": torch.tensor(sample.tirads, dtype=torch.long),
            "extension": torch.tensor(sample.extension, dtype=torch.long),
            "location": torch.tensor(sample.location, dtype=torch.long),
            "t_clinsig": torch.tensor(sample.t_clinsig, dtype=torch.long),
        }
